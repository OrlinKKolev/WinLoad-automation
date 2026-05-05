from __future__ import annotations

import sys
from pathlib import Path
import shutil

import openpyxl

project_root = Path(__file__).resolve().parent.parent.parent.parent
sys.path.insert(0, str(project_root))

import pandas as pd

from app.utils.numbers import parse_number


DATA_ROOT = project_root / "data"
SOURCE_EXCEL = DATA_ROOT / "schemes_output.xlsx"
TEMPLATE_EXCEL = DATA_ROOT / "STxxxx_Wind load check.xlsx"
REPORTS_DIR = DATA_ROOT / "reports"

CELL_MAP = [
    ("T11", "vb0", None),
    ("T18", "H", None),
    ("T19", "H", None),
    ("T20", "B", None),
    ("T21", "L", None),
    ("P25", "Elevation_m", None),
    ("P26", "ce_z", None),
    ("P27", None, 1.0),
]


def populate_scheme(row: dict, template_path: Path, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb.active

    ws.protection.sheet = False

    for cell_addr, col_name, fixed_value in CELL_MAP:
        if fixed_value is not None:
            ws[cell_addr] = fixed_value
        else:
            raw = row.get(col_name)
            if raw is None:
                raise ValueError(f"Column '{col_name}' not found in source row")
            ws[cell_addr] = parse_number(raw)

    ws.protection.sheet = False
    wb.save(str(output_path))
    print(f"  Saved: {output_path}")


def export_active_sheet_page1_to_pdf(excel_path: Path, pdf_path: Path) -> None:
    import win32com.client

    excel = win32com.client.Dispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.ScreenUpdating = False

    wb = None
    try:
        wb = excel.Workbooks.Open(str(excel_path.resolve()))
        ws = wb.ActiveSheet

        pdf_path.parent.mkdir(parents=True, exist_ok=True)

        ws.ExportAsFixedFormat(
            Type=0,
            Filename=str(pdf_path.resolve()),
            Quality=0,
            IncludeDocProperties=True,
            IgnorePrintAreas=False,
            From=1,
            To=1,
            OpenAfterPublish=False,
        )
    finally:
        if wb is not None:
            wb.Close(SaveChanges=False)
        excel.Quit()


def main() -> None:
    if not SOURCE_EXCEL.exists():
        raise FileNotFoundError(f"Source Excel not found: {SOURCE_EXCEL}")
    if not TEMPLATE_EXCEL.exists():
        raise FileNotFoundError(f"Template Excel not found: {TEMPLATE_EXCEL}")

    df = pd.read_excel(SOURCE_EXCEL, sheet_name="Schemes")
    if df.empty:
        raise RuntimeError("Schemes sheet is empty — nothing to process.")

    required = ["SchemeID", "H", "B", "L", "Elevation_m", "ce_z"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise RuntimeError(f"Source Excel is missing columns: {missing}")

    errors = 0
    for _, row in df.iterrows():
        scheme_id = row["SchemeID"]

        if str(row.get("ce_z", "")).upper() in ("ERROR", "NONE", "NAN", ""):
            print(f"[{scheme_id}] SKIPPED — ce_z is not a valid result")
            continue

        print(f"[{scheme_id}] Populating wind check Excel...")
        output_path = REPORTS_DIR / f"scheme_{scheme_id}_Wind_load_check.xlsx"
        pdf_path = REPORTS_DIR / f"scheme_{scheme_id}_Wind_load_check.pdf"

        try:
            populate_scheme(row.to_dict(), TEMPLATE_EXCEL, output_path)
            export_active_sheet_page1_to_pdf(output_path, pdf_path)
            print(f"[{scheme_id}] Done.")
        except Exception as e:
            print(f"[{scheme_id}] ERROR: {e}")
            errors += 1

    print(f"\\nAll schemes processed. {errors} error(s).")


if __name__ == "__main__":
    main()