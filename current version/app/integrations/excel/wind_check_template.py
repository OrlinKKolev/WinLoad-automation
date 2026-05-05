from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

from app.utils.numbers import parse_number


CELL_MAP = [
    ("T11", "vb0", None),
    ("T18", "H", None),
    ("T19", "H", None),
    ("T20", "B", None),
    ("T21", "L", None),
    ("P25", "Elevation_m", None),
    ("P26", "ce_z", None),
    ("P27", None, 1.0),
]


def populate_wind_check_template(row: dict, template_path: Path, output_path: Path) -> None:
    """Copy the template, populate mapped cells, and save the filled workbook."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb.active

    ws.protection.sheet = False

    for cell_addr, col_name, fixed_value in CELL_MAP:
        if fixed_value is not None:
            ws[cell_addr] = fixed_value
        else:
            raw = row.get(col_name)
            if raw is None:
                raise ValueError(f"Column '{col_name}' not found in source row")
            ws[cell_addr] = parse_number(raw)

    ws.protection.sheet = False
    wb.save(str(output_path))
    print(f"  Saved: {output_path}")